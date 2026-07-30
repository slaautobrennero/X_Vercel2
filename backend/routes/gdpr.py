"""
routes/gdpr.py
Endpoint per compliance GDPR:
- Impostazioni Sindacato (dati titolare del trattamento)
- Documenti legali (privacy, cookie, termini) renderizzati da template + settings,
  con possibilità di override tramite editor SuperAdmin (step-up 2FA)
- Diritti dell'interessato: accesso dati, esportazione (JSON/ZIP), cancellazione
  con grace period 30 gg + anonimizzazione post-31/03 anno successivo
- Notifica non bloccante nuove versioni documenti
- Export registro trattamenti (CSV/XLSX/PDF) per audit esterni
"""
import csv
import io
import os
import zipfile
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Optional

from bson import ObjectId
from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from core.audit import _log_audit
from core.auth import get_current_user
from core.db import db
from core.legal_templates import (
    PRIVACY_TEMPLATE, COOKIE_TEMPLATE, TERMINI_TEMPLATE,
    PRIVACY_VERSION, COOKIE_VERSION, TERMINI_VERSION,
    REGISTRO_TRATTAMENTI,
)
from core.roles import user_has_role
from core.totp import verify_code as verify_totp_code

router = APIRouter()

# Mapping tipo → (template, versione di default)
LEGAL_TEMPLATES = {
    "privacy": (PRIVACY_TEMPLATE, PRIVACY_VERSION),
    "cookie": (COOKIE_TEMPLATE, COOKIE_VERSION),
    "termini": (TERMINI_TEMPLATE, TERMINI_VERSION),
}

# ==================== MODELS ====================

class ImpostazioniSindacatoIn(BaseModel):
    denominazione: str = "SLA - CISAL"
    sede_legale: str = "Via Arpinati 22/1 - 16035 Rapallo GE"
    codice_fiscale: str = "91027960102"
    partita_iva: Optional[str] = None
    pec: str = ""
    email_privacy: str = ""
    sito_web: str = ""
    dpo_nome: Optional[str] = None
    dpo_email: Optional[str] = None
    foro_competente: str = "Genova"


DEFAULT_SETTINGS = {
    "denominazione": "SLA - CISAL",
    "sede_legale": "Via Arpinati 22/1 - 16035 Rapallo GE",
    "codice_fiscale": "91027960102",
    "partita_iva": None,
    "pec": "sla-cisal@pec.example.it",  # placeholder, modificabile
    "email_privacy": "privacy@sla-cisal.example.it",  # placeholder
    "sito_web": "https://www.sla-cisal.example.it",  # placeholder
    "dpo_nome": None,
    "dpo_email": None,
    "foro_competente": "Genova",
}


# ==================== HELPERS ====================

async def get_settings() -> dict:
    """Restituisce le impostazioni sindacato correnti (o default se mancanti)."""
    doc = await db.impostazioni_sindacato.find_one({"_id": "singleton"})
    if not doc:
        # Prima chiamata: crea default
        doc = {"_id": "singleton", **DEFAULT_SETTINGS}
        await db.impostazioni_sindacato.insert_one(doc)
    return {k: v for k, v in doc.items() if k != "_id"}


def render_template(template: str, settings: dict) -> str:
    """Sostituisce i placeholder {{...}} nel template con i valori dai settings."""
    oggi = datetime.now().strftime("%d/%m/%Y")
    piva_riga = f"Partita IVA: {settings['partita_iva']}\n" if settings.get("partita_iva") else ""
    sito_web_riga = f"Sito web: {settings['sito_web']}\n" if settings.get("sito_web") else ""
    dpo_info = "Non è stato nominato un DPO obbligatorio ai sensi dell'art. 37 GDPR."
    if settings.get("dpo_nome"):
        dpo_email = settings.get("dpo_email", "-")
        dpo_info = f"**DPO nominato**: {settings['dpo_nome']} — email: {dpo_email}"

    replacements = {
        "{{denominazione}}": settings["denominazione"],
        "{{sede_legale}}": settings["sede_legale"],
        "{{codice_fiscale}}": settings["codice_fiscale"],
        "{{piva_riga}}": piva_riga,
        "{{pec}}": settings["pec"],
        "{{email_privacy}}": settings["email_privacy"],
        "{{sito_web_riga}}": sito_web_riga,
        "{{dpo_info}}": dpo_info,
        "{{foro_competente}}": settings["foro_competente"],
        "{{oggi}}": oggi,
        "{{privacy_version}}": PRIVACY_VERSION,
        "{{cookie_version}}": COOKIE_VERSION,
        "{{termini_version}}": TERMINI_VERSION,
    }
    out = template
    for k, v in replacements.items():
        out = out.replace(k, str(v))
    return out


# ==================== IMPOSTAZIONI SINDACATO ====================

@router.get("/impostazioni-sindacato")
async def get_impostazioni_sindacato():
    """Pubblico: dati titolare del trattamento (usati anche nell'informativa)."""
    return await get_settings()


@router.put("/impostazioni-sindacato")
async def update_impostazioni_sindacato(payload: ImpostazioniSindacatoIn, request: Request):
    """Solo SuperAdmin: aggiorna dati sindacato."""
    user = await get_current_user(request)
    if not user_has_role(user, "superadmin"):
        raise HTTPException(status_code=403, detail="Solo SuperAdmin")
    data = payload.model_dump()
    await db.impostazioni_sindacato.update_one(
        {"_id": "singleton"}, {"$set": data}, upsert=True
    )
    await _log_audit(user, "impostazioni.update", target_type="impostazioni", note=str(data))
    return {"ok": True, "settings": data}


# ==================== DOCUMENTI LEGALI (pubblici) ====================
# Se in DB `documenti_legali` esiste una versione custom → usa quella,
# altrimenti renderizza il template hardcoded con i settings correnti.

async def _get_doc(tipo: str) -> dict:
    """Restituisce {version, contenuto} per un documento legale."""
    template, default_version = LEGAL_TEMPLATES[tipo]
    settings = await get_settings()
    doc = await db.documenti_legali.find_one({"_id": tipo})
    if doc:
        return {
            "version": doc.get("version", default_version),
            "contenuto": doc.get("contenuto", render_template(template, settings)),
            "updated_at": doc.get("updated_at"),
            "updated_by": doc.get("updated_by_nome"),
            "custom": True,
        }
    return {
        "version": default_version,
        "contenuto": render_template(template, settings),
        "updated_at": None,
        "updated_by": None,
        "custom": False,
    }


@router.get("/privacy")
async def get_privacy():
    return await _get_doc("privacy")


@router.get("/cookie-policy")
async def get_cookie_policy():
    return await _get_doc("cookie")


@router.get("/termini")
async def get_termini():
    return await _get_doc("termini")


@router.get("/versioni-documenti")
async def get_versioni_documenti():
    """Versioni correnti dei documenti legali (utile per re-consenso/notifica)."""
    return {
        "privacy_version": (await _get_doc("privacy"))["version"],
        "cookie_version": (await _get_doc("cookie"))["version"],
        "termini_version": (await _get_doc("termini"))["version"],
    }


# ==================== EDITOR DOCUMENTI LEGALI (SuperAdmin) ====================

class DocumentoLegaleIn(BaseModel):
    contenuto: str
    version: Optional[str] = None
    totp_code: Optional[str] = None  # step-up 2FA se abilitato


@router.get("/documenti-legali/{tipo}/anteprima")
async def anteprima_documento(tipo: str, request: Request):
    """SuperAdmin: anteprima del template renderizzato con settings correnti (non salva)."""
    if tipo not in LEGAL_TEMPLATES:
        raise HTTPException(status_code=404, detail="Tipo documento non valido")
    user = await get_current_user(request)
    if not user_has_role(user, "superadmin"):
        raise HTTPException(status_code=403, detail="Solo SuperAdmin")
    template, default_version = LEGAL_TEMPLATES[tipo]
    settings = await get_settings()
    return {
        "tipo": tipo,
        "version": default_version,
        "contenuto": render_template(template, settings),
        "template_originale": template,
    }


@router.put("/documenti-legali/{tipo}")
async def update_documento_legale(tipo: str, payload: DocumentoLegaleIn, request: Request):
    """
    SuperAdmin: sovrascrive/aggiorna il documento legale con contenuto custom.
    Richiede step-up 2FA se l'utente ha TOTP abilitato.
    Incrementa automaticamente la versione se non specificata.
    """
    if tipo not in LEGAL_TEMPLATES:
        raise HTTPException(status_code=404, detail="Tipo documento non valido")
    user = await get_current_user(request)
    if not user_has_role(user, "superadmin"):
        raise HTTPException(status_code=403, detail="Solo SuperAdmin")

    # Step-up 2FA se abilitato per questo utente
    if user.get("totp_enabled"):
        if not payload.totp_code:
            raise HTTPException(status_code=401, detail="Codice 2FA richiesto per modificare documenti legali")
        if not verify_totp_code(user.get("totp_secret", ""), payload.totp_code):
            raise HTTPException(status_code=401, detail="Codice 2FA non valido")

    # Calcolo prossima versione (se non fornita)
    existing = await db.documenti_legali.find_one({"_id": tipo})
    if payload.version:
        new_version = payload.version.strip()
    else:
        current = existing.get("version") if existing else LEGAL_TEMPLATES[tipo][1]
        # Incremento minor: "1.0" → "1.1", "1.5" → "1.6", non numerico → append ".1"
        try:
            major, minor = current.split(".", 1)
            new_version = f"{major}.{int(minor) + 1}"
        except (ValueError, AttributeError):
            new_version = f"{current}.1"

    now = datetime.now(timezone.utc)
    doc = {
        "_id": tipo,
        "contenuto": payload.contenuto,
        "version": new_version,
        "updated_at": now.isoformat(),
        "updated_by": user["id"],
        "updated_by_nome": f"{user.get('nome', '')} {user.get('cognome', '')}".strip(),
    }
    await db.documenti_legali.replace_one({"_id": tipo}, doc, upsert=True)

    # Storico versioni
    await db.documenti_legali_storico.insert_one({
        **doc,
        "_id": ObjectId(),
        "tipo": tipo,
    })

    await _log_audit(
        user, f"gdpr.documento.{tipo}.update",
        target_type="documento_legale", target_id=tipo,
        note=f"nuova versione: {new_version}",
    )
    return {"ok": True, "version": new_version, "updated_at": now.isoformat()}


# ==================== NOTIFICA NON BLOCCANTE NUOVI DOCUMENTI ====================

@router.get("/gdpr/nuovi-documenti")
async def check_nuovi_documenti(request: Request):
    """
    Ritorna quali documenti legali sono stati aggiornati DOPO l'ultima accettazione utente.
    Frontend mostra banner non bloccante di semplice notifica.
    """
    user = await get_current_user(request)
    user_oid = ObjectId(user["id"])
    user_doc = await db.users.find_one({"_id": user_oid})
    visti = (user_doc or {}).get("documenti_visti", {})

    updates = []
    for tipo in ("privacy", "cookie", "termini"):
        doc = await _get_doc(tipo)
        current_v = doc["version"]
        seen_v = visti.get(tipo)
        if seen_v != current_v:
            updates.append({
                "tipo": tipo,
                "version_corrente": current_v,
                "version_vista": seen_v,
                "updated_at": doc.get("updated_at"),
            })
    return {"nuovi": updates, "totale": len(updates)}


@router.post("/gdpr/marca-documenti-visti")
async def marca_documenti_visti(request: Request):
    """Marca tutti i documenti legali correnti come 'visti' dall'utente (dismiss banner)."""
    user = await get_current_user(request)
    user_oid = ObjectId(user["id"])
    versioni = {
        "privacy": (await _get_doc("privacy"))["version"],
        "cookie": (await _get_doc("cookie"))["version"],
        "termini": (await _get_doc("termini"))["version"],
    }
    await db.users.update_one(
        {"_id": user_oid},
        {"$set": {"documenti_visti": versioni, "documenti_visti_at": datetime.now(timezone.utc).isoformat()}},
    )
    return {"ok": True, "versioni": versioni}


# ==================== DIRITTI DELL'INTERESSATO ====================

@router.get("/gdpr/my-data")
async def get_my_data(request: Request):
    """Diritto di accesso (art. 15 GDPR): tutti i dati personali dell'utente."""
    user = await get_current_user(request)
    user_id_str = user["id"]
    user_oid = ObjectId(user_id_str)

    # Dati anagrafici
    anagrafica = {k: v for k, v in user.items() if k not in ("password_hash", "totp_secret")}
    if "sede_id" in anagrafica and anagrafica["sede_id"]:
        anagrafica["sede_id"] = str(anagrafica["sede_id"])

    # Rimborsi
    rimborsi_cursor = db.rimborsi.find({"user_id": user_oid})
    rimborsi = []
    async for r in rimborsi_cursor:
        r["_id"] = str(r["_id"])
        r["user_id"] = str(r["user_id"])
        if r.get("sede_id"):
            r["sede_id"] = str(r["sede_id"])
        rimborsi.append(r)

    # Consensi
    consensi_cursor = db.consensi_privacy.find({"user_id": user_oid})
    consensi = []
    async for c in consensi_cursor:
        c["_id"] = str(c["_id"])
        c["user_id"] = str(c["user_id"])
        consensi.append(c)

    return {
        "anagrafica": anagrafica,
        "rimborsi": rimborsi,
        "consensi_privacy": consensi,
        "estratto_il": datetime.now(timezone.utc).isoformat(),
        "info": "Questo estratto include tutti i tuoi dati personali gestiti dal portale.",
    }


class RichiestaCancellazioneIn(BaseModel):
    motivo: Optional[str] = None


@router.post("/gdpr/richiedi-cancellazione")
async def richiedi_cancellazione(payload: RichiestaCancellazioneIn, request: Request):
    """
    Diritto all'oblio (art. 17 GDPR): richiesta cancellazione con grace period 30gg.
    - Dopo 30gg dall'accettazione: account viene DISABILITATO (login bloccato)
    - Dopo 31/03 dell'anno successivo alla richiesta: dati anagrafici anonimizzati
      (i rimborsi contabili restano per obbligo fiscale)
    """
    user = await get_current_user(request)
    user_oid = ObjectId(user["id"])

    # Verifica se esiste già una richiesta pending
    existing = await db.richieste_cancellazione.find_one({
        "user_id": user_oid,
        "stato": {"$in": ["pending", "disabilitato"]},
    })
    if existing:
        raise HTTPException(
            status_code=400,
            detail="Hai già una richiesta di cancellazione in corso. Puoi annullarla dalla stessa sezione.",
        )

    now = datetime.now(timezone.utc)
    data_disabilitazione = now + timedelta(days=30)
    # Anonimizzazione: 31/03 dell'anno successivo alla richiesta (obbligo contabile)
    anno_purge = now.year + 1
    data_anonimizzazione = datetime(anno_purge, 3, 31, 23, 59, 59, tzinfo=timezone.utc)
    doc = {
        "user_id": user_oid,
        "email": user["email"],
        "motivo": payload.motivo or "",
        "stato": "pending",
        "richiesta_il": now,
        "esecuzione_prevista": data_disabilitazione,
        "anonimizzazione_prevista": data_anonimizzazione,
        "ip": request.client.host if request.client else None,
    }
    result = await db.richieste_cancellazione.insert_one(doc)
    await _log_audit(user, "gdpr.richiesta_cancellazione", target_type="user", target_id=user["id"], note=payload.motivo or "")
    return {
        "ok": True,
        "id": str(result.inserted_id),
        "esecuzione_prevista": data_disabilitazione.isoformat(),
        "anonimizzazione_prevista": data_anonimizzazione.isoformat(),
        "messaggio": (
            "Richiesta ricevuta. Hai 30 giorni per annullarla. "
            f"Trascorsi i quali il tuo account sarà disabilitato. "
            f"I dati anagrafici saranno anonimizzati dopo il {data_anonimizzazione.strftime('%d/%m/%Y')} "
            "(i rimborsi contabili sono conservati per obbligo fiscale)."
        ),
    }


@router.post("/gdpr/annulla-cancellazione")
async def annulla_cancellazione(request: Request):
    user = await get_current_user(request)
    user_oid = ObjectId(user["id"])
    result = await db.richieste_cancellazione.update_one(
        {"user_id": user_oid, "stato": "pending"},
        {"$set": {"stato": "annullata", "annullata_il": datetime.now(timezone.utc)}},
    )
    if result.modified_count == 0:
        raise HTTPException(status_code=404, detail="Nessuna richiesta in corso da annullare.")
    await _log_audit(user, "gdpr.annulla_cancellazione", target_type="user", target_id=user["id"])
    return {"ok": True, "messaggio": "Richiesta annullata. Il tuo account resta attivo."}


@router.get("/gdpr/stato-cancellazione")
async def stato_cancellazione(request: Request):
    """Stato attuale della richiesta di cancellazione dell'utente loggato."""
    user = await get_current_user(request)
    user_oid = ObjectId(user["id"])
    doc = await db.richieste_cancellazione.find_one(
        {"user_id": user_oid, "stato": {"$in": ["pending", "disabilitato"]}}
    )
    if not doc:
        return {"pending": False}
    # MongoDB restituisce datetime naive: rendo entrambi tz-aware per il diff
    esec = doc["esecuzione_prevista"]
    if esec.tzinfo is None:
        esec = esec.replace(tzinfo=timezone.utc)
    richiesta = doc["richiesta_il"]
    if richiesta.tzinfo is None:
        richiesta = richiesta.replace(tzinfo=timezone.utc)
    anon = doc.get("anonimizzazione_prevista")
    if anon and anon.tzinfo is None:
        anon = anon.replace(tzinfo=timezone.utc)
    return {
        "pending": True,
        "stato": doc.get("stato", "pending"),
        "richiesta_il": richiesta.isoformat(),
        "esecuzione_prevista": esec.isoformat(),
        "anonimizzazione_prevista": anon.isoformat() if anon else None,
        "giorni_residui": (esec - datetime.now(timezone.utc)).days,
    }


# ==================== EXPORT ZIP DATI UTENTE (art. 20 GDPR portabilità) ====================

UPLOAD_DIRS = [
    Path("/app/backend/uploads"),
    Path("/app/uploads"),
    Path(os.environ.get("UPLOAD_DIR", "/app/backend/uploads")),
]


def _find_upload_file(filename: str) -> Optional[Path]:
    """Cerca un file di ricevuta/contabile nelle directory di upload conosciute."""
    if not filename:
        return None
    for d in UPLOAD_DIRS:
        p = d / filename
        if p.exists() and p.is_file():
            return p
    return None


@router.get("/gdpr/export-zip")
async def export_zip(request: Request):
    """
    Diritto alla portabilità (art. 20 GDPR): archivio ZIP con dati JSON + ricevute/contabili.
    """
    import json
    user = await get_current_user(request)
    user_oid = ObjectId(user["id"])

    # Dati JSON completi
    anagrafica = {k: v for k, v in user.items() if k not in ("password_hash", "totp_secret")}
    if "sede_id" in anagrafica and anagrafica["sede_id"]:
        anagrafica["sede_id"] = str(anagrafica["sede_id"])

    rimborsi = []
    file_da_includere = set()
    async for r in db.rimborsi.find({"user_id": user_oid}):
        r["_id"] = str(r["_id"])
        r["user_id"] = str(r["user_id"])
        if r.get("sede_id"):
            r["sede_id"] = str(r["sede_id"])
        # Colleziona file allegati
        for rc in r.get("ricevute", []) or []:
            if isinstance(rc, dict) and rc.get("filename"):
                file_da_includere.add(rc["filename"])
        if r.get("contabile_filename"):
            file_da_includere.add(r["contabile_filename"])
        rimborsi.append(r)

    consensi = []
    async for c in db.consensi_privacy.find({"user_id": user_oid}):
        c["_id"] = str(c["_id"])
        c["user_id"] = str(c["user_id"])
        consensi.append(c)

    payload = {
        "anagrafica": anagrafica,
        "rimborsi": rimborsi,
        "consensi_privacy": consensi,
        "estratto_il": datetime.now(timezone.utc).isoformat(),
        "info": "Estratto GDPR (art. 20 - portabilità). Contiene tutti i tuoi dati personali.",
    }

    # Crea ZIP in memoria
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("dati.json", json.dumps(payload, indent=2, ensure_ascii=False, default=str))
        # Aggiunge file allegati se presenti
        for fname in file_da_includere:
            p = _find_upload_file(fname)
            if p:
                zf.write(p, arcname=f"ricevute/{fname}")
        # Info privacy
        settings = await get_settings()
        zf.writestr("README.txt", (
            "Estratto dei tuoi dati personali dal Portale SLA.\n\n"
            f"Titolare: {settings['denominazione']}\n"
            f"Email privacy: {settings['email_privacy']}\n"
            f"PEC: {settings['pec']}\n\n"
            "Il file dati.json contiene tutti i tuoi dati in formato leggibile.\n"
            "La cartella ricevute/ contiene gli allegati caricati per i rimborsi.\n"
        ))

    buf.seek(0)
    await _log_audit(user, "gdpr.export_zip", target_type="user", target_id=user["id"])
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename=dati_{user['email']}_{datetime.now().strftime('%Y%m%d')}.zip"},
    )


# ==================== REGISTRO TRATTAMENTI (SuperAdmin) ====================

@router.get("/gdpr/registro-trattamenti")
async def get_registro_trattamenti(request: Request):
    """Solo SuperAdmin: registro trattamenti art. 30 GDPR."""
    user = await get_current_user(request)
    if not user_has_role(user, "superadmin"):
        raise HTTPException(status_code=403, detail="Solo SuperAdmin")
    return {"trattamenti": REGISTRO_TRATTAMENTI, "titolare": await get_settings()}


@router.get("/gdpr/registro-trattamenti/export")
async def export_registro_trattamenti(request: Request, formato: str = "csv"):
    """SuperAdmin: export registro trattamenti in CSV/XLSX/PDF per audit esterni."""
    user = await get_current_user(request)
    if not user_has_role(user, "superadmin"):
        raise HTTPException(status_code=403, detail="Solo SuperAdmin")

    settings = await get_settings()
    rows = REGISTRO_TRATTAMENTI
    headers = ["nome", "finalita", "base_giuridica", "categorie_interessati",
               "categorie_dati", "destinatari", "trasferimenti_extra_ue",
               "conservazione", "misure_sicurezza"]

    if formato == "csv":
        out = io.StringIO()
        writer = csv.DictWriter(out, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)
        return StreamingResponse(
            iter([out.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=registro_trattamenti.csv"},
        )
    elif formato == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Registro Trattamenti"
        # Intestazione titolare
        ws.append([f"Titolare: {settings['denominazione']} — CF: {settings['codice_fiscale']}"])
        ws.append([f"Registro trattamenti art. 30 GDPR — generato il {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
        ws.append([])
        header_fill = PatternFill(start_color="1E4D8C", end_color="1E4D8C", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        header_row = 4
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=header_row, column=i, value=h.replace("_", " ").title())
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        for r in rows:
            ws.append([r.get(h, "") for h in headers])
        for col in ws.columns:
            max_len = max((len(str(c.value)) for c in col if c.value), default=15)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return StreamingResponse(
            iter([out.getvalue()]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=registro_trattamenti.xlsx"},
        )
    elif formato == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        out = io.BytesIO()
        doc = SimpleDocTemplate(out, pagesize=landscape(A4),
                                 rightMargin=8*mm, leftMargin=8*mm, topMargin=12*mm, bottomMargin=12*mm)
        styles = getSampleStyleSheet()
        elements = [
            Paragraph(f"<b>Registro dei Trattamenti (art. 30 GDPR)</b>", styles["Title"]),
            Paragraph(f"Titolare: <b>{settings['denominazione']}</b> — CF {settings['codice_fiscale']}", styles["Normal"]),
            Paragraph(f"Generato il {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles["Normal"]),
            Spacer(1, 6*mm),
        ]
        # Riga per riga (tabelle troppo larghe non stanno in una pagina)
        for r in rows:
            elements.append(Paragraph(f"<b>{r['nome']}</b>", styles["Heading3"]))
            data = [[k.replace("_", " ").title(), str(r.get(k, ""))] for k in headers if k != "nome"]
            t = Table(data, colWidths=[45*mm, 220*mm])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (0, -1), colors.HexColor("#E8EEF7")),
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor("#CCCCCC")),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING', (0, 0), (-1, -1), 4),
                ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ]))
            elements.append(t)
            elements.append(Spacer(1, 4*mm))
        doc.build(elements)
        out.seek(0)
        return StreamingResponse(
            iter([out.getvalue()]),
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=registro_trattamenti.pdf"},
        )
    else:
        raise HTTPException(status_code=400, detail="Formato non supportato (usa csv|xlsx|pdf)")


@router.get("/gdpr/richieste-cancellazione")
async def list_richieste_cancellazione(request: Request):
    """Solo SuperAdmin: lista richieste cancellazione (pending + disabilitate)."""
    user = await get_current_user(request)
    if not user_has_role(user, "superadmin"):
        raise HTTPException(status_code=403, detail="Solo SuperAdmin")
    cursor = db.richieste_cancellazione.find(
        {"stato": {"$in": ["pending", "disabilitato"]}}
    ).sort("richiesta_il", -1)
    out = []
    async for r in cursor:
        r["_id"] = str(r["_id"])
        r["user_id"] = str(r["user_id"])
        out.append(r)
    return {"richieste": out, "totale": len(out)}

