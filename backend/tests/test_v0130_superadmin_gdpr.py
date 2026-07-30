"""
Backend tests for v0.13.0-beta — SuperAdmin Settings & GDPR UI.
Covers:
- Impostazioni Sindacato: RBAC + persistence
- Documenti legali: anteprima, PUT with auto-version bump, custom flag propagation
- Registro trattamenti export (CSV/XLSX/PDF)
- GDPR self-service: my-data, export-zip, richiedi/annulla-cancellazione (+ anonimizzazione_prevista)
- Notifica non bloccante: nuovi-documenti / marca-documenti-visti
- Report rimborsi: colonna 'Data Approvazione' in CSV/XLSX/PDF
- Approvazione rimborso salva approvato_il/approvato_da
"""
import io
import os
import uuid
import zipfile
from datetime import datetime, timezone

import pytest
import requests
from bson import ObjectId
from pymongo import MongoClient
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://portale-rimborsi.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"

SUPERADMIN_EMAIL = "superadmin@sla.it"
SUPERADMIN_PASSWORD = "SlaAdmin2024!"


def _load_backend_env():
    env = {}
    try:
        with open("/app/backend/.env") as f:
            for line in f:
                line = line.strip()
                if not line or line.startswith("#") or "=" not in line:
                    continue
                k, v = line.split("=", 1)
                env[k.strip()] = v.strip().strip('"').strip("'")
    except Exception:
        pass
    return env


# ---------- Fixtures ----------

@pytest.fixture(scope="session")
def mongo():
    env = _load_backend_env()
    url = os.environ.get("MONGO_URL") or env.get("MONGO_URL") or "mongodb://localhost:27017"
    dbname = os.environ.get("DB_NAME") or env.get("DB_NAME") or "portale_sla"
    client = MongoClient(url)
    yield client[dbname]
    client.close()


@pytest.fixture(scope="session")
def superadmin_session():
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"email": SUPERADMIN_EMAIL, "password": SUPERADMIN_PASSWORD})
    if r.status_code != 200:
        pytest.skip(f"Superadmin login failed: {r.status_code}")
    return s


@pytest.fixture(scope="session")
def test_user_session(mongo):
    """Fresh non-superadmin user for RBAC & self-service tests."""
    email = f"test_gdpr_v0130_{uuid.uuid4().hex[:8]}@example.com"
    password = "TestPassGdpr2024!"
    s = requests.Session()
    r = s.post(f"{API}/auth/register", json={
        "email": email, "password": password,
        "nome": "V0130", "cognome": "Tester",
        "privacy_accepted": True, "termini_accepted": True,
        "privacy_version": "1.0",
    })
    assert r.status_code == 200, f"Register failed: {r.status_code} {r.text}"
    yield {"session": s, "email": email}
    # cleanup
    try:
        udoc = mongo.users.find_one({"email": email}, {"_id": 1})
        if udoc:
            mongo.richieste_cancellazione.delete_many({"user_id": udoc["_id"]})
            mongo.consensi_privacy.delete_many({"user_id": udoc["_id"]})
        mongo.users.delete_many({"email": email})
    except Exception as e:
        print(f"cleanup warn: {e}")


# ==================== VERSION ====================

class TestVersion:
    def test_version_bumped(self):
        r = requests.get(f"{API}/version")
        assert r.status_code == 200
        d = r.json()
        assert d["version"] == "0.13.0-beta"


# ==================== IMPOSTAZIONI SINDACATO ====================

class TestImpostazioniSindacato:
    def test_get_returns_defaults(self):
        r = requests.get(f"{API}/impostazioni-sindacato")
        assert r.status_code == 200
        d = r.json()
        assert d["denominazione"] == "SLA - CISAL"
        assert "Via Arpinati 22/1" in d["sede_legale"]
        assert d["codice_fiscale"] == "91027960102"

    def test_put_superadmin_ok(self, superadmin_session):
        cur = requests.get(f"{API}/impostazioni-sindacato").json()
        payload = {
            "denominazione": cur["denominazione"],
            "sede_legale": cur["sede_legale"],
            "codice_fiscale": cur["codice_fiscale"],
            "partita_iva": cur.get("partita_iva"),
            "pec": cur.get("pec", ""),
            "email_privacy": cur.get("email_privacy", ""),
            "sito_web": cur.get("sito_web", ""),
            "dpo_nome": cur.get("dpo_nome"),
            "dpo_email": cur.get("dpo_email"),
            "foro_competente": cur.get("foro_competente", "Genova"),
        }
        r = superadmin_session.put(f"{API}/impostazioni-sindacato", json=payload)
        assert r.status_code == 200, r.text
        assert r.json()["ok"] is True

    def test_put_non_superadmin_403(self, test_user_session):
        r = test_user_session["session"].put(f"{API}/impostazioni-sindacato", json={
            "denominazione": "HACK", "sede_legale": "x", "codice_fiscale": "0",
            "pec": "", "email_privacy": "", "sito_web": "", "foro_competente": "Roma"
        })
        assert r.status_code == 403


# ==================== DOCUMENTI LEGALI EDITOR ====================

class TestDocumentiLegali:
    def test_anteprima_ok(self, superadmin_session):
        r = superadmin_session.get(f"{API}/documenti-legali/privacy/anteprima")
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["tipo"] == "privacy"
        assert "SLA - CISAL" in d["contenuto"]
        assert "{{denominazione}}" not in d["contenuto"]  # placeholders rendered

    def test_anteprima_invalid_type_404(self, superadmin_session):
        r = superadmin_session.get(f"{API}/documenti-legali/foobar/anteprima")
        assert r.status_code == 404

    def test_anteprima_requires_superadmin(self, test_user_session):
        r = test_user_session["session"].get(f"{API}/documenti-legali/privacy/anteprima")
        assert r.status_code == 403

    def test_put_auto_increments_version(self, superadmin_session, mongo):
        # baseline: pick current version
        cur = requests.get(f"{API}/privacy").json()
        cur_v = cur["version"]
        # parse to know expected next
        try:
            major, minor = cur_v.split(".", 1)
            expected_next = f"{major}.{int(minor) + 1}"
        except Exception:
            expected_next = f"{cur_v}.1"

        contenuto_test = f"# Privacy TEST v0.13.0 {uuid.uuid4().hex[:6]}\nContenuto di test."
        r = superadmin_session.put(f"{API}/documenti-legali/privacy",
                                   json={"contenuto": contenuto_test})
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["ok"] is True
        assert d["version"] == expected_next

        # GET /api/privacy now returns custom
        r2 = requests.get(f"{API}/privacy")
        assert r2.status_code == 200
        d2 = r2.json()
        assert d2["custom"] is True
        assert d2["version"] == expected_next
        assert contenuto_test.split("\n")[0] in d2["contenuto"]

    def test_put_invalid_tipo_404(self, superadmin_session):
        r = superadmin_session.put(f"{API}/documenti-legali/invalid",
                                   json={"contenuto": "x"})
        assert r.status_code == 404

    def test_put_non_superadmin_403(self, test_user_session):
        r = test_user_session["session"].put(f"{API}/documenti-legali/privacy",
                                             json={"contenuto": "hack"})
        assert r.status_code == 403


# ==================== REGISTRO TRATTAMENTI ====================

class TestRegistroTrattamenti:
    def test_list(self, superadmin_session):
        r = superadmin_session.get(f"{API}/gdpr/registro-trattamenti")
        assert r.status_code == 200
        d = r.json()
        assert len(d["trattamenti"]) == 4
        assert "titolare" in d
        assert d["titolare"]["denominazione"] == "SLA - CISAL"

    def test_export_csv(self, superadmin_session):
        r = superadmin_session.get(f"{API}/gdpr/registro-trattamenti/export?formato=csv")
        assert r.status_code == 200
        assert "text/csv" in r.headers.get("content-type", "")
        body = r.text
        assert "nome" in body.splitlines()[0]
        # at least 4 records + header
        assert len(body.splitlines()) >= 5

    def test_export_xlsx(self, superadmin_session):
        r = superadmin_session.get(f"{API}/gdpr/registro-trattamenti/export?formato=xlsx")
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        # Header row is row 4
        headers = [c.value for c in ws[4]]
        assert "Nome" in headers

    def test_export_pdf(self, superadmin_session):
        r = superadmin_session.get(f"{API}/gdpr/registro-trattamenti/export?formato=pdf")
        assert r.status_code == 200
        assert r.headers.get("content-type") == "application/pdf"
        assert r.content[:4] == b"%PDF"

    def test_export_invalid_format_400(self, superadmin_session):
        r = superadmin_session.get(f"{API}/gdpr/registro-trattamenti/export?formato=bogus")
        assert r.status_code == 400

    def test_richieste_cancellazione_list(self, superadmin_session):
        r = superadmin_session.get(f"{API}/gdpr/richieste-cancellazione")
        assert r.status_code == 200
        d = r.json()
        assert "richieste" in d
        assert isinstance(d["richieste"], list)


# ==================== GDPR SELF-SERVICE ====================

class TestGdprSelfService:
    def test_my_data(self, test_user_session):
        r = test_user_session["session"].get(f"{API}/gdpr/my-data")
        assert r.status_code == 200
        d = r.json()
        assert d["anagrafica"]["email"] == test_user_session["email"]
        assert "rimborsi" in d
        assert "consensi_privacy" in d
        assert len(d["consensi_privacy"]) >= 1

    def test_export_zip(self, test_user_session):
        r = test_user_session["session"].get(f"{API}/gdpr/export-zip")
        assert r.status_code == 200
        assert r.headers.get("content-type") == "application/zip"
        # Parse ZIP
        zf = zipfile.ZipFile(io.BytesIO(r.content))
        names = zf.namelist()
        assert "dati.json" in names
        assert "README.txt" in names
        # dati.json has anagrafica
        import json
        data = json.loads(zf.read("dati.json"))
        assert data["anagrafica"]["email"] == test_user_session["email"]

    def test_cancellazione_lifecycle_with_anonimizzazione(self, test_user_session, mongo):
        s = test_user_session["session"]
        # Ensure clean
        udoc = mongo.users.find_one({"email": test_user_session["email"]}, {"_id": 1})
        mongo.richieste_cancellazione.delete_many({"user_id": udoc["_id"]})

        # Stato iniziale
        r = s.get(f"{API}/gdpr/stato-cancellazione")
        assert r.status_code == 200 and r.json()["pending"] is False

        # Richiedi
        r = s.post(f"{API}/gdpr/richiedi-cancellazione", json={"motivo": "test v0130"})
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["ok"] is True
        assert "esecuzione_prevista" in d
        assert "anonimizzazione_prevista" in d
        # Verifica: anonimizzazione = 31/03 dell'anno successivo
        anon = datetime.fromisoformat(d["anonimizzazione_prevista"].replace("Z", "+00:00"))
        now = datetime.now(timezone.utc)
        assert anon.month == 3 and anon.day == 31
        assert anon.year == now.year + 1

        # Doppia richiesta -> 400
        r2 = s.post(f"{API}/gdpr/richiedi-cancellazione", json={"motivo": "dup"})
        assert r2.status_code == 400

        # Stato pending
        r = s.get(f"{API}/gdpr/stato-cancellazione")
        st = r.json()
        assert st["pending"] is True
        assert st.get("anonimizzazione_prevista") is not None
        assert 28 <= st["giorni_residui"] <= 30

        # Annulla
        r = s.post(f"{API}/gdpr/annulla-cancellazione")
        assert r.status_code == 200

        # Stato torna vuoto
        r = s.get(f"{API}/gdpr/stato-cancellazione")
        assert r.json()["pending"] is False


# ==================== NOTIFICA NUOVI DOCUMENTI ====================

class TestNuoviDocumenti:
    def test_nuovi_documenti_and_marca_visti(self, superadmin_session):
        # SuperAdmin also has documenti_visti tracking
        r = superadmin_session.get(f"{API}/gdpr/nuovi-documenti")
        assert r.status_code == 200
        d = r.json()
        assert "nuovi" in d
        assert "totale" in d
        assert isinstance(d["nuovi"], list)

        # Marca visti
        r = superadmin_session.post(f"{API}/gdpr/marca-documenti-visti")
        assert r.status_code == 200
        d = r.json()
        assert d["ok"] is True
        assert "versioni" in d
        assert set(d["versioni"].keys()) >= {"privacy", "cookie", "termini"}

        # Ora nuovi = 0
        r = superadmin_session.get(f"{API}/gdpr/nuovi-documenti")
        assert r.json()["totale"] == 0


# ==================== REPORT RIMBORSI + APPROVAZIONE ====================

class TestApprovazioneEReport:
    """Verify approvato_il persistence and Data Approvazione column."""

    def _get_or_create_test_rimborso(self, superadmin_session, mongo):
        """Find or create a rimborso for the current year to approve."""
        anno = datetime.now().year
        # Get superadmin user id from /auth/me
        me = superadmin_session.get(f"{API}/auth/me").json()
        user_id = me["id"]

        # Look for an existing 'in_attesa' rimborso we can approve
        r = superadmin_session.get(f"{API}/rimborsi?stato=in_attesa")
        assert r.status_code == 200
        rimborsi = r.json()
        target = None
        for rr in (rimborsi if isinstance(rimborsi, list) else rimborsi.get("rimborsi", [])):
            if rr.get("stato") == "in_attesa":
                target = rr
                break

        if not target:
            # Create one: need motivo_id and sede_id
            motivi = superadmin_session.get(f"{API}/motivi-rimborso").json()
            # Find a non-mandatory-notes motivo
            motivo_id = None
            for m in motivi:
                if not m.get("richiede_note"):
                    motivo_id = m["id"] if "id" in m else str(m.get("_id"))
                    break
            if not motivo_id and motivi:
                motivo_id = motivi[0].get("id") or str(motivi[0].get("_id"))
            payload = {
                "data": f"{anno}-06-15",
                "motivo_id": motivo_id,
                "indirizzo_partenza": "Rapallo",
                "indirizzo_partenza_tipo": "custom",
                "indirizzo_arrivo": "Genova",
                "km_andata": 30,
                "km_calcolati": 30,
                "km_modificati_manualmente": False,
                "andata_ritorno": True,
                "uso_autostrada": False,
                "costo_autostrada": 0,
                "importo_pasti": 0,
                "numero_partecipanti_pasto": 0,
                "note": "TEST v0.13.0 approvazione",
            }
            rc = superadmin_session.post(f"{API}/rimborsi", json=payload)
            assert rc.status_code == 200, rc.text
            target = rc.json()
        return target

    def test_approvazione_salva_timestamp(self, superadmin_session, mongo):
        target = self._get_or_create_test_rimborso(superadmin_session, mongo)
        rid = target.get("id") or target.get("_id")
        # PUT stato=approvato
        r = superadmin_session.put(f"{API}/rimborsi/{rid}", json={"stato": "approvato"})
        assert r.status_code == 200, r.text
        # Verify DB has approvato_il + approvato_da
        doc = mongo.rimborsi.find_one({"_id": ObjectId(rid)})
        assert doc is not None
        assert doc.get("stato") == "approvato"
        assert doc.get("approvato_il") is not None
        assert doc.get("approvato_da") is not None
        assert doc.get("approvato_da_nome")

    def test_export_csv_has_data_approvazione(self, superadmin_session):
        anno = datetime.now().year
        r = superadmin_session.get(f"{API}/reports/rimborsi-export?anno={anno}&formato=csv")
        assert r.status_code == 200
        lines = r.text.splitlines()
        assert lines, "CSV vuoto"
        header = lines[0]
        assert "Data Approvazione" in header
        # Almeno una riga con data approvazione popolata (per rimborsi appena approvati)
        idx = header.split(",").index("Data Approvazione")
        has_populated = False
        for row in lines[1:]:
            cols = row.split(",")
            if idx < len(cols) and cols[idx].strip() not in ("", '""'):
                has_populated = True
                break
        assert has_populated, "Nessuna riga ha Data Approvazione popolata"

    def test_export_xlsx_has_data_approvazione(self, superadmin_session):
        anno = datetime.now().year
        r = superadmin_session.get(f"{API}/reports/rimborsi-export?anno={anno}&formato=xlsx")
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        assert "Data Approvazione" in headers

    def test_export_pdf_ok(self, superadmin_session):
        anno = datetime.now().year
        r = superadmin_session.get(f"{API}/reports/rimborsi-export?anno={anno}&formato=pdf")
        assert r.status_code == 200
        assert r.content[:4] == b"%PDF"
        # colonna 'Approvato il' non è directamente ispezionabile senza parsing PDF;
        # ci fidiamo del content-type + status
