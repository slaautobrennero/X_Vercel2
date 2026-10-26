# v0.14.3 - Fix audit di sicurezza: SEC-001 permessi utenti, revoca sessioni (token_version),
# must_change_password, CSRF Origin, IP reale per lockout, upload magic-bytes, link annunci,
# export CSV/XLSX anti-formula, notifiche altrui, avvisi-sicurezza superadmin
import csv
import io
import random
import uuid
from datetime import UTC, datetime

import bcrypt
import pytest
import requests
from bson import ObjectId
from dotenv import dotenv_values
from pymongo import MongoClient

from tests.conftest import cartella_test

FE_ENV = dotenv_values("/app/frontend/.env")
BE_ENV = dotenv_values("/app/backend/.env")
BASE_URL = FE_ENV["REACT_APP_BACKEND_URL"].rstrip("/")
API = f"{BASE_URL}/api"
LOCAL_API = "http://localhost:8001/api"  # CF-Connecting-IP viene bloccato da Cloudflare sull'URL pubblico
PWD = "Prova123!"
SA_EMAIL, SA_PWD = "superadmin@sla.it", "SlaAdmin2024!"
EVIL = "https://evil.example.com"

mc = MongoClient(BE_ENV["MONGO_URL"])
db = mc[BE_ENV["DB_NAME"]]
HASH = bcrypt.hashpw(PWD.encode(), bcrypt.gensalt()).decode()
TAG = uuid.uuid4().hex[:6]

PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n"
PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
FAKE = b"<html><script>alert(1)</script></html>"


def _rand_ip():
    return f"10.{random.randint(0, 255)}.{random.randint(0, 255)}.{random.randint(1, 254)}"


def _login(email, pwd=PWD):
    s = requests.Session()
    s.headers["X-Forwarded-For"] = _rand_ip()  # evita rate limit 10/min condiviso
    r = s.post(f"{API}/auth/login", json={"email": email, "password": pwd})
    return s, r


def _mk_user(key, ruoli, sede_id, **extra):
    email = f"test_{key}_{TAG}@test.it".lower()
    doc = {"email": email, "password_hash": HASH, "nome": "TEST", "cognome": key, "ruoli": ruoli,
           "sede_id": sede_id, "disabled": False, "iban": "IT60X0542811101000000123456",
           "created_at": datetime.now(UTC).isoformat(), **extra}
    return email, str(db.users.insert_one(doc).inserted_id)


@pytest.fixture(scope="module")
def env():
    sa, r = _login(SA_EMAIL, SA_PWD)
    assert r.status_code == 200, r.text
    sa_id = r.json()["id"]
    sA = str(db.sedi.insert_one({"nome": f"TEST_SedeA_{TAG}", "codice": f"TA{TAG}", "tariffa_km": 0.35,
                                 "rimborso_pasti": 15.0, "rimborso_autostrada": True}).inserted_id)
    sB = str(db.sedi.insert_one({"nome": f"TEST_SedeB_{TAG}", "codice": f"TB{TAG}", "tariffa_km": 0.35,
                                 "rimborso_pasti": 15.0, "rimborso_autostrada": True}).inserted_id)
    u = {}
    for key, ruoli, sede in [("admA", ["admin"], sA), ("admA2", ["admin"], sA), ("delA", ["delegato"], sA),
                             ("delA2", ["delegato"], sA), ("delA3", ["delegato"], sA), ("delA4", ["delegato"], sA),
                             ("segA", ["segretario"], sA), ("suA", ["superuser"], sA), ("delB", ["delegato"], sB)]:
        u[key] = _mk_user(key, ruoli, sede)
    yield {"sa": sa, "sa_id": sa_id, "sA": sA, "sB": sB, "u": u}
    ids = [ObjectId(v[1]) for v in u.values()]
    uid_strs = [v[1] for v in u.values()]
    db.users.delete_many({"_id": {"$in": ids}})
    db.sedi.delete_many({"_id": {"$in": [ObjectId(sA), ObjectId(sB)]}})
    db.notifiche.delete_many({"user_id": {"$in": uid_strs}})
    db.rimborsi.delete_many({"user_id": {"$in": uid_strs}})
    db.annunci.delete_many({"titolo": {"$regex": f"^TEST_{TAG}"}})
    db.documenti.delete_many({"nome": {"$regex": f"^TEST_{TAG}"}})
    db.audit_log.delete_many({"target_id": {"$in": uid_strs}})
    db.login_attempts.delete_many({"identifier": {"$regex": TAG}})


def _s(env, key):
    s, r = _login(env["u"][key][0])
    assert r.status_code == 200, r.text
    return s


# ---------- SEC-001 PUT /users/{id} ----------
class TestUpdateUser:
    def test_admin_other_sede_403(self, env):
        s = _s(env, "admA")
        r = s.put(f"{API}/users/{env['u']['delB'][1]}", json={"nome": "Hack"})
        assert r.status_code == 403, r.text
        assert db.users.find_one({"_id": ObjectId(env['u']['delB'][1])})["nome"] == "TEST"

    def test_admin_on_other_admin_403(self, env):
        s = _s(env, "admA")
        r = s.put(f"{API}/users/{env['u']['admA2'][1]}", json={"iban": "IT00HACK"})
        assert r.status_code == 403, r.text
        r = s.put(f"{API}/users/{env['u']['suA'][1]}", json={"iban": "IT00HACK"})
        assert r.status_code == 403, r.text

    def test_admin_on_own_delegato_200_audit(self, env):
        s = _s(env, "admA")
        uid = env['u']['delA'][1]
        r = s.put(f"{API}/users/{uid}", json={"iban": "IT11NEWIBAN0000000000000001"})
        assert r.status_code == 200, r.text
        assert db.users.find_one({"_id": ObjectId(uid)})["iban"] == "IT11NEWIBAN0000000000000001"
        a = db.audit_log.find_one({"action": "user.update_profile", "target_id": uid}, sort=[("created_at", -1)])
        assert a is not None
        assert "IT60X0542811101000000123456" in a["old_value"]
        assert "IT11NEWIBAN0000000000000001" in a["new_value"]

    def test_self_update_200(self, env):
        s = _s(env, "delB")
        r = s.put(f"{API}/users/{env['u']['delB'][1]}", json={"telefono": "3331234567"})
        assert r.status_code == 200, r.text
        assert db.users.find_one({"_id": ObjectId(env['u']['delB'][1])})["telefono"] == "3331234567"

    def test_delegato_on_other_403(self, env):
        s = _s(env, "delB")
        r = s.put(f"{API}/users/{env['u']['delA'][1]}", json={"nome": "x"})
        assert r.status_code == 403


# ---------- SEC-001 PUT /users/{id}/ruolo ----------
class TestUpdateRole:
    @pytest.mark.parametrize("ruolo", ["admin", "superuser", "superadmin"])
    def test_admin_cannot_assign_privileged(self, env, ruolo):
        s = _s(env, "admA")
        r = s.put(f"{API}/users/{env['u']['delA4'][1]}/ruolo", json={"ruoli": [ruolo]})
        assert r.status_code == 400, r.text
        assert db.users.find_one({"_id": ObjectId(env['u']['delA4'][1])})["ruoli"] == ["delegato"]

    def test_admin_other_sede_403(self, env):
        s = _s(env, "admA")
        r = s.put(f"{API}/users/{env['u']['delB'][1]}/ruolo", json={"ruoli": ["cassiere"]})
        assert r.status_code == 403, r.text

    def test_admin_on_other_admin_403(self, env):
        s = _s(env, "admA")
        r = s.put(f"{API}/users/{env['u']['admA2'][1]}/ruolo", json={"ruoli": ["iscritto"]})
        assert r.status_code == 403, r.text

    def test_admin_own_roles_403(self, env):
        s = _s(env, "admA")
        r = s.put(f"{API}/users/{env['u']['admA'][1]}/ruolo", json={"ruoli": ["delegato"]})
        assert r.status_code == 403, r.text

    def test_admin_on_own_delegato_ok(self, env):
        s = _s(env, "admA")
        r = s.put(f"{API}/users/{env['u']['delA4'][1]}/ruolo", json={"ruoli": ["delegato", "cassiere"]})
        assert r.status_code == 200, r.text
        assert set(db.users.find_one({"_id": ObjectId(env['u']['delA4'][1])})["ruoli"]) == {"delegato", "cassiere"}

    def test_superadmin_can_all(self, env):
        uid = env['u']['delB'][1]
        r = env["sa"].put(f"{API}/users/{uid}/ruolo", json={"ruoli": ["admin", "delegato"]})
        assert r.status_code == 200, r.text
        r = env["sa"].put(f"{API}/users/{uid}/ruolo", json={"ruoli": ["delegato"]})
        assert r.status_code == 200
        assert db.users.find_one({"_id": ObjectId(uid)})["ruoli"] == ["delegato"]


# ---------- reset-password / toggle-disabled ----------
class TestResetToggle:
    @pytest.mark.parametrize("actor,target", [("segA", "admA"), ("segA", "suA"), ("segA", "delB"),
                                              ("admA", "admA2"), ("admA", "delB")])
    def test_reset_forbidden(self, env, actor, target):
        s = _s(env, actor)
        r = s.post(f"{API}/users/{env['u'][target][1]}/reset-password")
        assert r.status_code == 403, r.text
        r = s.put(f"{API}/users/{env['u'][target][1]}/toggle-disabled", json={"disabled": True})
        assert r.status_code == 403, r.text
        assert not db.users.find_one({"_id": ObjectId(env['u'][target][1])}).get("disabled")

    def test_segretario_on_superadmin_403(self, env):
        s = _s(env, "segA")
        r = s.post(f"{API}/users/{env['sa_id']}/reset-password")
        assert r.status_code == 403


# ---------- Revoca sessioni + must_change_password ----------
class TestSessionRevocation:
    def test_change_password_revokes(self, env):
        email, uid = env['u']['delA2']
        s = _s(env, "delA2")
        old = dict(s.cookies)
        r = s.post(f"{API}/auth/change-password", json={"current_password": PWD, "new_password": "Nuova123!x"})
        assert r.status_code == 200, r.text
        o = requests.Session()
        o.cookies.update(old)
        assert o.get(f"{API}/auth/me").status_code == 401
        assert o.post(f"{API}/auth/refresh").status_code == 401
        db.users.update_one({"_id": ObjectId(uid)}, {"$set": {"password_hash": HASH}})

    def test_disable_revokes(self, env):
        email, uid = env['u']['delA3']
        s = _s(env, "delA3")
        adm = _s(env, "admA")
        r = adm.put(f"{API}/users/{uid}/toggle-disabled", json={"disabled": True})
        assert r.status_code == 200, r.text
        assert s.get(f"{API}/auth/me").status_code == 401
        r = adm.put(f"{API}/users/{uid}/toggle-disabled", json={"disabled": False})
        assert r.status_code == 200
        # dopo riattivazione il vecchio token resta revocato (tv incrementata)
        assert s.get(f"{API}/auth/me").status_code == 401
        assert s.post(f"{API}/auth/refresh").status_code == 401

    def test_reset_revokes_and_must_change(self, env):
        email, uid = env['u']['delA3']
        s = _s(env, "delA3")
        adm = _s(env, "admA")
        r = adm.post(f"{API}/users/{uid}/reset-password")
        assert r.status_code == 200, r.text
        temp = r.json()["temporary_password"]
        assert s.get(f"{API}/auth/me").status_code == 401
        assert s.post(f"{API}/auth/refresh").status_code == 401
        n, r = _login(email, temp)
        assert r.status_code == 200, r.text
        me = n.get(f"{API}/auth/me")
        assert me.status_code == 200 and me.json().get("must_change_password") is True
        r = n.get(f"{API}/rimborsi")
        assert r.status_code == 403 and r.json()["detail"] == "PASSWORD_CHANGE_REQUIRED"
        assert n.get(f"{API}/notifiche").status_code == 200
        r = n.post(f"{API}/auth/change-password", json={"current_password": temp, "new_password": "Nuova123!x"})
        assert r.status_code == 200, r.text
        n2, r = _login(email, "Nuova123!x")
        assert r.status_code == 200
        assert n2.get(f"{API}/rimborsi").status_code == 200
        db.users.update_one({"_id": ObjectId(uid)}, {"$set": {"password_hash": HASH}})


# ---------- CSRF Origin ----------
class TestCSRF:
    def test_evil_origin_blocked(self, env):
        s = _s(env, "delB")
        for m, path, body in [("post", "/auth/logout", None), ("put", "/notifiche/letto-tutte", None),
                              ("put", f"/users/{env['u']['delB'][1]}", {"nome": "evil"}),
                              ("delete", "/annunci/000000000000000000000000", None)]:
            r = getattr(s, m)(f"{API}{path}", json=body, headers={"Origin": EVIL})
            assert r.status_code == 403, f"{m} {path}: {r.status_code}"
            assert r.json()["detail"] == "Origine della richiesta non consentita"
        r = s.put(f"{API}/users/{env['u']['delB'][1]}", json={"nome": "evil"}, headers={"Referer": EVIL + "/x"})
        assert r.status_code == 403
        assert db.users.find_one({"_id": ObjectId(env['u']['delB'][1])})["nome"] == "TEST"

    def test_same_origin_ok(self, env):
        s = _s(env, "delB")
        r = s.put(f"{API}/notifiche/letto-tutte", headers={"Origin": BASE_URL})
        assert r.status_code == 200, r.text


# ---------- Lockout per IP reale ----------
class TestLockoutIP:
    def test_cf_ip_separate_counters(self, env):
        email = f"test_lock_{TAG}@test.it"
        codes = []
        for _ in range(6):
            r = requests.post(f"{LOCAL_API}/auth/login", json={"email": email, "password": "bad"},
                              headers={"CF-Connecting-IP": "1.1.1.1"})
            codes.append(r.status_code)
        assert codes[:5] == [401] * 5 and codes[5] == 429, codes
        r = requests.post(f"{LOCAL_API}/auth/login", json={"email": email, "password": "bad"},
                          headers={"CF-Connecting-IP": "2.2.2.2"})
        assert r.status_code == 401
        db.login_attempts.delete_many({"identifier": {"$regex": email}})


# ---------- Upload magic bytes ----------
class TestUpload:
    def test_documenti(self, env):
        s = _s(env, "admA")
        cid = cartella_test(env["sA"])
        r = s.post(f"{API}/documenti", files={"file": ("a.pdf", FAKE, "application/pdf")},
                   data={"nome": f"TEST_{TAG}_fake", "cartella_id": cid})
        assert r.status_code == 400, r.text
        r = s.post(f"{API}/documenti", files={"file": ("a.pdf", PDF, "application/pdf")},
                   data={"nome": f"TEST_{TAG}_ok", "cartella_id": cid})
        assert r.status_code == 200, r.text
        r = s.post(f"{API}/documenti", files={"file": ("img.pdf", PNG, "application/pdf")},
                   data={"nome": f"TEST_{TAG}_png", "cartella_id": cid})
        assert r.status_code == 200, r.text
        d = db.documenti.find_one({"nome": f"TEST_{TAG}_png"})
        assert d["path"].endswith(".png")

    def test_ricevute(self, env):
        s = _s(env, "delA")
        motivo = db.motivi_rimborso.find_one({"nome": "RSA"})
        r = s.post(f"{API}/rimborsi", json={"data": "2099-01-15", "motivo_id": str(motivo["_id"]),
                                            "indirizzo_partenza": "A", "indirizzo_arrivo": "B", "km_andata": 10,
                                            "note": "=SUM(1)"})
        assert r.status_code == 200, r.text
        rid = r.json().get("id")
        env["rid"] = rid
        r = s.post(f"{API}/rimborsi/{rid}/ricevute", files={"file": ("r.pdf", FAKE, "application/pdf")})
        assert r.status_code == 400, r.text
        r = s.post(f"{API}/rimborsi/{rid}/ricevute", files={"file": ("r.jpg", b"\xff\xd8\xff\xe0" + b"0" * 32, "image/jpeg")})
        assert r.status_code == 200, r.text
        assert r.json()["path"].endswith(".jpg")

    def test_annunci_upload_and_link(self, env):
        s = _s(env, "admA")
        r = s.post(f"{API}/annunci", data={"titolo": f"TEST_{TAG}_f", "contenuto": "x"},
                   files={"file": ("a.pdf", FAKE, "application/pdf")})
        assert r.status_code == 400, r.text
        r = s.post(f"{API}/annunci", data={"titolo": f"TEST_{TAG}_js", "contenuto": "x",
                                           "link_documento": "javascript:alert(1)"})
        assert r.status_code == 400, r.text
        r = s.post(f"{API}/annunci", data={"titolo": f"TEST_{TAG}_js2", "contenuto": "x",
                                           "link_documento": " JavaScript:alert(1)"})
        assert r.status_code == 400, r.text
        r = s.post(f"{API}/annunci", data={"titolo": f"TEST_{TAG}_ok", "contenuto": "x",
                                           "link_documento": "https://example.com/doc.pdf"},
                   files={"file": ("a.pdf", PDF, "application/pdf")})
        assert r.status_code == 200, r.text


# ---------- Export anti formula ----------
class TestExport:
    def test_csv_xlsx(self, env):
        sa = env["sa"]
        r = sa.get(f"{API}/reports/rimborsi-export", params={"anno": 2099, "formato": "csv"})
        assert r.status_code == 200, r.text
        rows = list(csv.DictReader(io.StringIO(r.text)))
        mine = [x for x in rows if x.get("Email") == env['u']['delA'][0]]
        assert mine, r.text[:500]
        assert mine[0]["Note"] == "'=SUM(1)"
        assert mine[0]["Totale"] == "7.00"
        r = sa.get(f"{API}/reports/rimborsi-export", params={"anno": 2099, "formato": "xlsx"})
        assert r.status_code == 200
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(r.content)).active
        vals = [c.value for row in ws.iter_rows() for c in row]
        assert "'=SUM(1)" in vals
        assert not any(isinstance(v, str) and v.startswith("=") for v in vals)

    def test_safe_cell_unit(self):
        import sys
        sys.path.insert(0, "/app/backend")
        from routes.reports import _safe_cell
        assert _safe_cell("-12.50") == "-12.50"
        assert _safe_cell("+3") == "+3"
        assert _safe_cell("@cmd") == "'@cmd"
        assert _safe_cell("-2+3+cmd|' /C calc'!A0") .startswith("'")
        assert _safe_cell(-5) == -5


# ---------- Notifiche altrui ----------
class TestNotifiche:
    def test_mark_other_users_notifica(self, env):
        nid = db.notifiche.insert_one({"user_id": env['u']['delB'][1], "titolo": "TEST", "messaggio": "x",
                                       "letto": False, "created_at": datetime.now(UTC).isoformat()}).inserted_id
        s = _s(env, "delA")
        r = s.put(f"{API}/notifiche/{nid}/letto")
        assert r.status_code in (200, 404)
        assert db.notifiche.find_one({"_id": nid})["letto"] is False
        s2 = _s(env, "delB")
        assert s2.put(f"{API}/notifiche/{nid}/letto").status_code == 200
        assert db.notifiche.find_one({"_id": nid})["letto"] is True


# ---------- Avvisi sicurezza ----------
class TestAvvisi:
    def test_superadmin_only(self, env):
        r = env["sa"].get(f"{API}/system/avvisi-sicurezza")
        assert r.status_code == 200
        assert isinstance(r.json()["avvisi"], list)
        assert any("password iniziale" in a for a in r.json()["avvisi"])
        for k in ("admA", "suA", "delA"):
            assert _s(env, k).get(f"{API}/system/avvisi-sicurezza").status_code == 403
